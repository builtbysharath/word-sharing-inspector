"""Optional independent fixture generator. Requires python-docx >= 1.2."""
from pathlib import Path
from docx import Document

document = Document()
document.core_properties.author = "Taylor Example"
document.core_properties.last_modified_by = "Jordan Example"
document.core_properties.title = "Independent producer sample"
document.add_heading("Proposal for review", 0)
paragraph = document.add_paragraph("Delivery estimate: four weeks.")
document.add_comment(paragraph.runs, author="Taylor Example", initials="TE", text="Confirm this estimate before sharing.")
hidden = document.add_paragraph().add_run("Internal planning note, not intended for the recipient.")
hidden.font.hidden = True
document.sections[0].header.paragraphs[0].text = "Synthetic example — no client data"
document.save(Path(__file__).with_name("producer-sample.docx"))
